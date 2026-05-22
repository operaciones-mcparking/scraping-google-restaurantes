import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def main() -> int:
    if len(sys.argv) != 3:
        print("Uso: convert_csv_to_xlsx.py entrada.csv salida.xlsx")
        return 2

    csv_path = Path(sys.argv[1])
    xlsx_path = Path(sys.argv[2])

    wb = Workbook()
    ws = wb.active
    ws.title = "Restaurantes"

    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        for row in reader:
            ws.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        max_len = 0
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(column)].width = min(max(max_len + 2, 12), 42)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"XLSX creado: {xlsx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

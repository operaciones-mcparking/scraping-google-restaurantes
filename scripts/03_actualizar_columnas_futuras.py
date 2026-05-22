import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


FIELDS = [
    "Nombre restaurante",
    "Rating",
    "Cantidad reviews",
    "Dirección",
    "Teléfono",
    "Sitio web",
    "Categoría",
    "Google Maps URL",
    "Latitud",
    "Longitud",
    "Comuna",
    "Región",
    "País",
    "Fuente",
    "Fecha extracción",
    "Calidad dato",
    "Es restaurante válido",
    "Instagram URL",
    "Instagram Usuario",
    "Facebook URL",
    "TikTok URL",
    "Tiene Redes",
    "Calidad Redes",
    "Está en Uber Eats",
    "URL Uber Eats",
    "Confianza Uber Eats",
    "Está en PedidosYa",
    "URL PedidosYa",
    "Confianza PedidosYa",
    "Está en Rappi",
    "URL Rappi",
    "Confianza Rappi",
    "Observaciones",
]


def read_csv_rows(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def normalize_rows(rows: list[dict]) -> list[dict]:
    return [{field: row.get(field, "") for field in FIELDS} for row in rows]


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Restaurantes"
    ws.append(FIELDS)
    for row in rows:
        ws.append([row.get(field, "") for field in FIELDS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 42)

    wb.save(path)


def update_raw_json(path: Path, rows: list[dict]) -> None:
    if not path.exists():
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["rows"] = rows
    payload["schema"] = FIELDS
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def validate_xlsx(path: Path) -> None:
    ws = load_workbook(path).active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    missing = [field for field in FIELDS if field not in headers]
    if missing:
        raise RuntimeError(f"Faltan columnas en {path.name}: {missing}")


def update_dataset(base: Path) -> None:
    csv_path = base.with_suffix(".csv")
    xlsx_path = base.with_suffix(".xlsx")
    raw_path = base.parent / f"{base.name}_raw.json"

    rows = normalize_rows(read_csv_rows(csv_path))
    write_csv(csv_path, rows)
    write_xlsx(xlsx_path, rows)
    update_raw_json(raw_path, rows)
    validate_xlsx(xlsx_path)
    print(f"Actualizado: {base.name} ({len(rows)} filas, {len(FIELDS)} columnas)")


def main() -> int:
    if len(sys.argv) < 2:
        print("Uso: 03_actualizar_columnas_futuras.py base_sin_extension [base_sin_extension...]")
        return 2

    for arg in sys.argv[1:]:
        update_dataset(Path(arg))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

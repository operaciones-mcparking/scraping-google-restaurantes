import csv
import json
import urllib.parse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "enriquecimiento_redes_sociales.json"

SOCIAL_FIELDS = [
    "Tiene Instagram",
    "Calidad Instagram",
    "Tiene Facebook",
    "Calidad Facebook",
    "Búsqueda Instagram sugerida",
    "Búsqueda Facebook sugerida",
    "Estado enriquecimiento redes",
    "Observación redes",
]


def project_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def load_config() -> dict:
    return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))


def read_rows(config: dict) -> tuple[list[str], list[dict]]:
    wb = load_workbook(project_path(config["entradaExcel"]))
    ws = wb[config["hojaEntrada"]]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {headers[c - 1]: ws.cell(r, c).value or "" for c in range(1, ws.max_column + 1)}
        if row.get("Nivel comercial") == config["filtroNivelComercial"]:
            rows.append(row)
    rows.sort(key=lambda row: (int(row.get("Score comercial") or 0), int(str(row.get("Cantidad reviews") or "0").replace(".", "").replace(",", "") or 0)), reverse=True)
    return headers, rows


def search_url(query: str) -> str:
    return "https://www.google.com/search?q=" + urllib.parse.quote_plus(query)


def enrich_structure(row: dict) -> dict:
    name = row.get("Nombre restaurante", "")
    comuna = row.get("Comuna", "")
    region = row.get("Región", "")
    site = row.get("Sitio web", "")

    row = dict(row)
    row.setdefault("Instagram URL", "")
    row.setdefault("Instagram Usuario", "")
    row.setdefault("Facebook URL", "")
    row["Tiene Instagram"] = ""
    row["Calidad Instagram"] = ""
    row["Tiene Facebook"] = ""
    row["Calidad Facebook"] = ""
    row["Búsqueda Instagram sugerida"] = search_url(f'{name} {comuna} Instagram')
    row["Búsqueda Facebook sugerida"] = search_url(f'{name} {comuna} Facebook')
    row["Estado enriquecimiento redes"] = "Pendiente"
    row["Observación redes"] = f"Validar contra nombre, comuna y sitio web: {site}" if site else "Validar contra nombre y comuna; no hay sitio web."
    return row


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        writer.writerows(rows)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 55)


def write_xlsx(path: Path, fields: list[str], rows: list[dict], summary: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Restaurantes"
    ws.append(fields)
    for row in rows:
        ws.append([row.get(field, "") for field in fields])
    style_sheet(ws)

    summary_ws = wb.create_sheet("Resumen")
    summary_fields = ["Métrica", "Valor"]
    summary_ws.append(summary_fields)
    for row in summary:
        summary_ws.append([row["Métrica"], row["Valor"]])
    style_sheet(summary_ws)
    wb.save(path)


def summary_rows(total_high: int, pilot_count: int) -> list[dict]:
    return [
        {"Métrica": "Restaurantes Alto potencial disponibles", "Valor": total_high},
        {"Métrica": "Restaurantes en piloto", "Valor": pilot_count},
        {"Métrica": "Instagram/Facebook buscados automáticamente", "Valor": 0},
        {"Métrica": "Estado", "Valor": "Preparado, sin búsquedas externas"},
    ]


def main() -> int:
    config = load_config()
    base_fields, high_rows = read_rows(config)
    output_fields = base_fields + [field for field in SOCIAL_FIELDS if field not in base_fields]
    prepared = [enrich_structure(row) for row in high_rows]
    pilot = prepared[: int(config["maxPiloto"])]
    summary = summary_rows(len(prepared), len(pilot))

    write_csv(project_path(config["salidaCandidatosCsv"]), output_fields, prepared)
    write_xlsx(project_path(config["salidaCandidatosExcel"]), output_fields, prepared, summary)
    write_csv(project_path(config["salidaPilotoCsv"]), output_fields, pilot)
    write_xlsx(project_path(config["salidaPilotoExcel"]), output_fields, pilot, summary)
    project_path(config["salidaPilotoJson"]).write_text(json.dumps({
        "config": config,
        "totalAltoPotencial": len(prepared),
        "piloto": pilot,
        "nota": "Archivo preparado sin ejecutar búsquedas en Instagram o Facebook.",
    }, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Alto potencial disponibles: {len(prepared)}")
    print(f"Piloto preparado: {len(pilot)}")
    print(f"Excel piloto: {project_path(config['salidaPilotoExcel'])}")
    print(f"CSV piloto: {project_path(config['salidaPilotoCsv'])}")
    print(f"JSON piloto: {project_path(config['salidaPilotoJson'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

from __future__ import annotations

import argparse
import csv
import logging
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from supabase_client import get_supabase_client


MAX_LIMIT = 30
CANDIDATE_POOL = 300

OUTPUT_XLSX = ROOT / "data" / "rappi_piloto_revision_manual_30.xlsx"
OUTPUT_CSV = ROOT / "data" / "rappi_piloto_revision_manual_30.csv"
LOG_PATH = ROOT / "data" / "logs" / "rappi_enrichment.log"

HEADERS = [
    "Restaurante",
    "Comuna",
    "Query Google sugerida",
    "URL busqueda Google",
    "Query Rappi sugerida",
    "Estado revision Rappi",
    "URL Rappi encontrada",
    "Observacion",
    "Revisado manualmente",
]

STATUS_OPTIONS = ["Dudoso", "Encontrado", "No encontrado", "Error"]
REVIEW_OPTIONS = ["No", "Si"]


def setup_logger() -> logging.Logger:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("rappi_manual_review")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    file_handler = logging.FileHandler(LOG_PATH, encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    return logger


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Genera planilla manual para validar presencia en Rappi.")
    parser.add_argument("--limit", type=int, default=30, help="Maximo de restaurantes para el piloto. Se limita a 30.")
    parser.add_argument("--output", type=Path, default=OUTPUT_XLSX, help="Ruta del Excel de salida.")
    parser.add_argument("--csv-output", type=Path, default=OUTPUT_CSV, help="Ruta del CSV de salida.")
    return parser.parse_args()


def fetch_restaurants(limit: int, candidate_pool: int = CANDIDATE_POOL) -> list[dict[str, Any]]:
    if limit <= 0:
        return []
    client = get_supabase_client()
    response = (
        client.table("restaurantes")
        .select(
            "crm_id,nombre_restaurante,comuna,nivel_comercial,tipo_negocio,"
            "telefono,sitio_web,google_maps_url,nombre_normalizado"
        )
        .order("fecha_carga", desc=True)
        .limit(max(limit, min(candidate_pool, max(limit * 10, 120))))
        .execute()
    )
    return list(response.data or [])


def normalize_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.split())


def has_value(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return bool(text and text not in {"nan", "none", "null", "-"})


def chain_key(item: dict[str, Any]) -> str:
    raw_name = item.get("nombre_normalizado") or item.get("nombre_restaurante") or ""
    tokens = normalize_text(raw_name).replace("&", " ").split()
    stopwords = {
        "restaurant",
        "restaurante",
        "restobar",
        "bar",
        "pub",
        "cafe",
        "cafeteria",
        "cocina",
        "comida",
        "de",
        "del",
        "la",
        "las",
        "el",
        "los",
        "y",
        "chile",
        "sucursal",
    }
    useful_tokens = [token for token in tokens if token not in stopwords]
    return " ".join(useful_tokens[:3]) or normalize_text(raw_name)


def restaurant_score(item: dict[str, Any]) -> int:
    score = 0
    if normalize_text(item.get("nivel_comercial")) == "alto potencial":
        score += 60
    elif normalize_text(item.get("nivel_comercial")) == "medio potencial":
        score += 25
    if has_value(item.get("telefono")):
        score += 20
    if has_value(item.get("sitio_web")):
        score += 15
    if has_value(item.get("google_maps_url")):
        score += 5
    return score


def select_representative_sample(restaurants: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    if limit <= 0:
        return []

    candidates = sorted(
        restaurants,
        key=lambda item: (
            restaurant_score(item),
            normalize_text(item.get("nivel_comercial")) == "alto potencial",
            has_value(item.get("telefono")) or has_value(item.get("sitio_web")),
        ),
        reverse=True,
    )

    selected: list[dict[str, Any]] = []
    selected_keys: set[str] = set()
    selected_chains: set[str] = set()
    comuna_counts: dict[str, int] = {}
    type_counts: dict[str, int] = {}

    def stable_key(item: dict[str, Any]) -> str:
        google_url = normalize_text(item.get("google_maps_url"))
        if google_url:
            return google_url
        return "|".join(
            [
                normalize_text(item.get("nombre_restaurante")),
                normalize_text(item.get("comuna")),
                normalize_text(item.get("tipo_negocio")),
            ]
        )

    def pick(allow_chain_repeat: bool) -> bool:
        best_idx: int | None = None
        best_rank: tuple[int, int, int, int] | None = None
        for idx, item in enumerate(candidates):
            key = stable_key(item)
            if key in selected_keys:
                continue
            chain = chain_key(item)
            if not allow_chain_repeat and chain in selected_chains:
                continue

            comuna = normalize_text(item.get("comuna")) or "sin comuna"
            business_type = normalize_text(item.get("tipo_negocio")) or "sin tipo"
            rank = (
                -comuna_counts.get(comuna, 0),
                -type_counts.get(business_type, 0),
                restaurant_score(item),
                -idx,
            )
            if best_rank is None or rank > best_rank:
                best_idx = idx
                best_rank = rank

        if best_idx is None:
            return False

        item = candidates[best_idx]
        key = stable_key(item)
        chain = chain_key(item)
        comuna = normalize_text(item.get("comuna")) or "sin comuna"
        business_type = normalize_text(item.get("tipo_negocio")) or "sin tipo"

        selected.append(item)
        selected_keys.add(key)
        selected_chains.add(chain)
        comuna_counts[comuna] = comuna_counts.get(comuna, 0) + 1
        type_counts[business_type] = type_counts.get(business_type, 0) + 1
        return True

    while len(selected) < limit and pick(allow_chain_repeat=False):
        pass
    while len(selected) < limit and pick(allow_chain_repeat=True):
        pass

    return selected[:limit]


def google_query(name: str, comuna: str) -> str:
    parts = [name]
    if comuna:
        parts.append(comuna)
    parts.append("Rappi")
    return " ".join(part for part in parts if part).strip()


def google_search_url(query: str) -> str:
    return f"https://www.google.com/search?q={quote_plus(query)}"


def rappi_query(name: str) -> str:
    return name.strip()


def rappi_search_url(query: str) -> str:
    return f"https://www.rappi.cl/search?query={quote_plus(query)}"


def build_manual_rows(restaurants: list[dict[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for item in restaurants:
        name = str(item.get("nombre_restaurante") or "").strip()
        comuna = str(item.get("comuna") or "").strip()
        g_query = google_query(name, comuna)
        r_query = rappi_query(name)
        rows.append(
            {
                "Restaurante": name,
                "Comuna": comuna,
                "Query Google sugerida": g_query,
                "URL busqueda Google": google_search_url(g_query),
                "Query Rappi sugerida": r_query,
                "Estado revision Rappi": "Dudoso",
                "URL Rappi encontrada": "",
                "Observacion": "Pendiente de revision manual. No marcar encontrado sin evidencia clara.",
                "Revisado manualmente": "No",
            }
        )
    return rows


def write_csv(rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Revision Rappi"

    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(color="FFFFFF", bold=True)
    link_font = Font(color="2563EB", underline="single")

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row.get(header, "") for header in HEADERS])

    for row_idx in range(2, len(rows) + 2):
        google_cell = ws.cell(row=row_idx, column=4)
        google_cell.hyperlink = google_cell.value
        google_cell.font = link_font
        google_cell.style = "Hyperlink"

        rappi_query_cell = ws.cell(row=row_idx, column=5)
        rappi_query_cell.hyperlink = rappi_search_url(str(rappi_query_cell.value or ""))
        rappi_query_cell.font = link_font
        rappi_query_cell.style = "Hyperlink"

        rappi_url_cell = ws.cell(row=row_idx, column=7)
        rappi_url_cell.font = link_font

    widths = {
        "A": 34,
        "B": 18,
        "C": 42,
        "D": 58,
        "E": 34,
        "F": 22,
        "G": 46,
        "H": 58,
        "I": 22,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows(min_row=2, max_row=max(2, len(rows) + 1), min_col=1, max_col=len(HEADERS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if rows:
        status_validation = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=False)
        reviewed_validation = DataValidation(type="list", formula1='"' + ",".join(REVIEW_OPTIONS) + '"', allow_blank=False)
        ws.add_data_validation(status_validation)
        ws.add_data_validation(reviewed_validation)
        status_validation.add(f"F2:F{len(rows) + 1}")
        reviewed_validation.add(f"I2:I{len(rows) + 1}")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(1, len(rows) + 1)}"
    ws.sheet_view.showGridLines = False
    for idx in range(1, len(rows) + 2):
        ws.row_dimensions[idx].height = 24 if idx == 1 else 42

    wb.save(path)


def main() -> None:
    args = parse_args()
    logger = setup_logger()
    limit = max(0, min(args.limit, MAX_LIMIT))
    logger.info("Inicio piloto manual Rappi | limit=%s", limit)

    candidates = fetch_restaurants(limit)
    restaurants = select_representative_sample(candidates, limit)
    rows = build_manual_rows(restaurants)
    write_xlsx(rows, args.output)
    write_csv(rows, args.csv_output)

    high_potential = sum(1 for item in restaurants if normalize_text(item.get("nivel_comercial")) == "alto potencial")
    comunas = len({normalize_text(item.get("comuna")) for item in restaurants if normalize_text(item.get("comuna"))})
    business_types = len(
        {normalize_text(item.get("tipo_negocio")) for item in restaurants if normalize_text(item.get("tipo_negocio"))}
    )

    print(f"Restaurantes incluidos: {len(rows)}")
    print(f"Alto potencial incluidos: {high_potential}")
    print(f"Comunas representadas: {comunas}")
    print(f"Tipos de negocio representados: {business_types}")
    print(f"Excel generado: {args.output}")
    print(f"CSV generado: {args.csv_output}")
    print("No se escribio nada en Supabase.")
    logger.info(
        "Planilla manual generada | candidatos=%s | filas=%s | alto_potencial=%s | comunas=%s | tipos=%s | xlsx=%s | csv=%s",
        len(candidates),
        len(rows),
        high_potential,
        comunas,
        business_types,
        args.output,
        args.csv_output,
    )


if __name__ == "__main__":
    main()

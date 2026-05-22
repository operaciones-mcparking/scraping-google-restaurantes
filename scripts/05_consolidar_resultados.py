import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]

FIELDS = [
    "Nombre restaurante",
    "Rating",
    "Cantidad reviews",
    "Dirección",
    "Teléfono",
    "Sitio web",
    "Categoría",
    "Google Maps URL",
    "Latitud",
    "Longitud",
    "Comuna",
    "Región",
    "País",
    "Fuente",
    "Fecha extracción",
    "Calidad dato",
    "Es restaurante válido",
    "Instagram URL",
    "Instagram Usuario",
    "Facebook URL",
    "TikTok URL",
    "Tiene Redes",
    "Calidad Redes",
    "Está en Uber Eats",
    "URL Uber Eats",
    "Confianza Uber Eats",
    "Está en PedidosYa",
    "URL PedidosYa",
    "Confianza PedidosYa",
    "Está en Rappi",
    "URL Rappi",
    "Confianza Rappi",
    "Observaciones",
]

SUMMARY_FIELDS = [
    "Grupo",
    "Comuna",
    "Región",
    "Filas",
    "Restaurantes válidos",
    "No válidos",
    "Teléfonos vacíos",
    "Direcciones vacías",
    "Sitios web vacíos",
]


def project_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def read_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [{field: row.get(field, "") for field in FIELDS} for row in csv.DictReader(handle)]


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        writer.writerows(rows)


def add_sheet(wb: Workbook, name: str, fields: list[str], rows: list[dict]) -> None:
    ws = wb.create_sheet(name)
    ws.append(fields)
    for row in rows:
        ws.append([row.get(field, "") for field in fields])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 42)


def summarize(item: dict, rows: list[dict]) -> dict:
    return {
        "Grupo": item["grupo"],
        "Comuna": item["comuna"],
        "Región": item["region"],
        "Filas": len(rows),
        "Restaurantes válidos": sum(1 for row in rows if row.get("Es restaurante válido") == "Sí"),
        "No válidos": sum(1 for row in rows if row.get("Es restaurante válido") == "No"),
        "Teléfonos vacíos": sum(1 for row in rows if not row.get("Teléfono")),
        "Direcciones vacías": sum(1 for row in rows if not row.get("Dirección")),
        "Sitios web vacíos": sum(1 for row in rows if not row.get("Sitio web")),
    }


def main() -> int:
    config_path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "configs" / "scraper_google_maps_multi_comuna.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))

    all_rows = []
    summary_rows = []
    for item in config["comunas"]:
        csv_path = ROOT / "data" / "comunas" / f"{item['slug']}_restaurantes_30.csv"
        rows = read_rows(csv_path)
        all_rows.extend(rows)
        summary_rows.append(summarize(item, rows))

    consolidated_csv = project_path(config["salidaConsolidadaCsv"])
    summary_csv = project_path(config["salidaResumenCsv"])
    consolidated_xlsx = project_path(config["salidaConsolidadaExcel"])

    write_csv(consolidated_csv, all_rows, FIELDS)
    write_csv(summary_csv, summary_rows, SUMMARY_FIELDS)

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Base consolidada", FIELDS, all_rows)
    add_sheet(wb, "Resumen", SUMMARY_FIELDS, summary_rows)
    consolidated_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(consolidated_xlsx)

    print(f"Consolidado CSV: {consolidated_csv}")
    print(f"Resumen CSV: {summary_csv}")
    print(f"Consolidado Excel: {consolidated_xlsx}")
    print(f"Filas consolidadas: {len(all_rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

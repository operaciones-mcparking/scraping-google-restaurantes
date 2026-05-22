import csv
import math
import re
import unicodedata
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = ROOT / "data" / "consolidado_restaurantes_priorizados.xlsx"
OUTPUT_XLSX = ROOT / "data" / "consolidado_restaurantes_priorizados_normalizado.xlsx"
OUTPUT_CSV = ROOT / "data" / "consolidado_restaurantes_priorizados_normalizado.csv"

BASE_SHEET = "Base consolidada"

ADDED_FIELDS = [
    "Nombre normalizado",
    "Posible cadena/franquicia",
    "Grupo cadena/franquicia",
    "Posible duplicado entre comunas",
    "Grupo duplicado",
    "Tipo negocio",
    "Score comercial",
    "Nivel comercial",
]

SUMMARY_FIELDS = ["Métrica", "Valor"]
TOP_FIELDS = ["Ranking", "Nombre restaurante", "Comuna", "Región", "Rating", "Cantidad reviews", "Tipo negocio", "Nivel comercial", "Score comercial"]
COUNT_FIELDS = ["Nombre", "Cantidad"]


def strip_accents(text: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", text or "") if unicodedata.category(ch) != "Mn")


def norm_text(text: str) -> str:
    text = strip_accents(str(text or "")).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def parse_float(value) -> float:
    if value in (None, ""):
        return 0.0
    text = str(value).replace(",", ".")
    match = re.search(r"\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else 0.0


def parse_int(value) -> int:
    if value in (None, ""):
        return 0
    text = str(value)
    digits = re.sub(r"[^0-9]", "", text)
    return int(digits) if digits else 0


def load_rows(path: Path) -> tuple[list[str], list[dict]]:
    wb = load_workbook(path)
    ws = wb[BASE_SHEET]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {headers[c - 1]: ws.cell(r, c).value or "" for c in range(1, ws.max_column + 1)}
        rows.append(row)
    return headers, rows


def normalize_name(name: str) -> str:
    text = norm_text(name)
    remove_terms = [
        "restaurante",
        "restaurant",
        "restoran",
        "restorant",
        "restobar",
        "sucursal",
        "local",
        "las condes",
        "providencia",
        "vitacura",
        "lo barnechea",
        "la reina",
        "nunoa",
        "ñuñoa",
        "la florida",
        "curico",
        "chillan",
        "chile",
    ]
    for term in remove_terms:
        text = re.sub(rf"\b{re.escape(norm_text(term))}\b", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text or norm_text(name)


def chain_key(name: str) -> str:
    normalized = normalize_name(name)
    tokens = normalized.split()
    if len(tokens) >= 2:
        return " ".join(tokens[:2])
    return normalized


def classify_type(row: dict) -> str:
    text = norm_text(f"{row.get('Nombre restaurante')} {row.get('Categoría')} {row.get('Observaciones')}")
    if any(term in text for term in ["dark kitchen", "delivery only", "cocina oculta"]):
        return "Dark Kitchen"
    if any(term in text for term in ["cafeteria", " cafe", "cafe ", "coffee"]):
        return "Cafetería"
    if any(term in text for term in ["fast food", "comida rapida", "hamburgueseria", "burger", "kfc", "mcdonald", "wendy", "subway"]):
        return "Fast Food"
    if any(term in text for term in ["bar", "pub", "cerveceria", "taberna"]):
        return "Bar/Pub"
    if any(term in text for term in ["pasteleria", "bakery", "panaderia", "bolleria"]):
        return "Bakery/Pastelería"
    if any(term in text for term in ["restaurante", "restaurant", "restoran", "bistro", "cocina", "trattoria", "parrilla", "sushi", "pizzeria"]):
        return "Restaurante"
    return "Otro"


def commercial_score(row: dict) -> int:
    rating = parse_float(row.get("Rating"))
    reviews = parse_int(row.get("Cantidad reviews"))

    rating_points = min(max(rating, 0), 5) / 5 * 35
    review_points = min(math.log10(reviews + 1) / math.log10(10001), 1) * 35
    website_points = 10 if str(row.get("Sitio web") or "").strip() else 0
    phone_points = 10 if str(row.get("Teléfono") or "").strip() else 0

    quality = str(row.get("Calidad dato") or "")
    if quality == "Completo":
        quality_points = 10
    elif quality == "Parcial":
        quality_points = 6
    elif quality in ("Sin teléfono", "Sin dirección"):
        quality_points = 2
    else:
        quality_points = 4

    if str(row.get("Es restaurante válido") or "") == "No":
        quality_points = max(quality_points - 4, 0)

    return round(rating_points + review_points + website_points + phone_points + quality_points)


def commercial_level(score: int, row: dict) -> str:
    if str(row.get("Es restaurante válido") or "") == "No":
        return "Bajo potencial"
    if score >= 78:
        return "Alto potencial"
    if score >= 55:
        return "Medio potencial"
    return "Bajo potencial"


def duplicate_key(row: dict) -> str:
    return f"{normalize_name(row.get('Nombre restaurante'))}|{norm_text(row.get('Dirección'))}"


def find_cross_commune_duplicates(rows: list[dict]) -> dict[int, str]:
    groups: dict[int, str] = {}
    by_name = defaultdict(list)
    for idx, row in enumerate(rows):
        by_name[chain_key(row.get("Nombre restaurante"))].append((idx, row))

    group_id = 1
    for candidates in by_name.values():
        if len(candidates) < 2:
            continue
        used = set()
        for i, (idx_a, row_a) in enumerate(candidates):
            if idx_a in used:
                continue
            group = [idx_a]
            name_a = normalize_name(row_a.get("Nombre restaurante"))
            for idx_b, row_b in candidates[i + 1:]:
                if row_a.get("Comuna") == row_b.get("Comuna"):
                    continue
                name_b = normalize_name(row_b.get("Nombre restaurante"))
                similarity = SequenceMatcher(None, name_a, name_b).ratio()
                if similarity >= 0.9:
                    group.append(idx_b)
            if len(group) > 1:
                label = f"DUP-{group_id:03d}"
                for idx in group:
                    groups[idx] = label
                    used.add(idx)
                group_id += 1
    return groups


def enrich_rows(rows: list[dict]) -> list[dict]:
    chain_counts = Counter(chain_key(row.get("Nombre restaurante")) for row in rows)
    duplicate_groups = find_cross_commune_duplicates(rows)

    enriched = []
    for idx, row in enumerate(rows):
        row = dict(row)
        key = chain_key(row.get("Nombre restaurante"))
        score = commercial_score(row)
        row["Nombre normalizado"] = normalize_name(row.get("Nombre restaurante"))
        row["Posible cadena/franquicia"] = "Sí" if chain_counts[key] >= 2 else "No"
        row["Grupo cadena/franquicia"] = key if chain_counts[key] >= 2 else ""
        row["Posible duplicado entre comunas"] = "Sí" if idx in duplicate_groups else "No"
        row["Grupo duplicado"] = duplicate_groups.get(idx, "")
        row["Tipo negocio"] = classify_type(row)
        row["Score comercial"] = score
        row["Nivel comercial"] = commercial_level(score, row)
        enriched.append(row)
    return enriched


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        writer.writerows(rows)


def add_sheet(wb: Workbook, name: str, fields: list[str], rows: list[dict]) -> None:
    ws = wb.create_sheet(name)
    ws.append(fields)
    for row in rows:
        ws.append([row.get(field, "") for field in fields])
    style_sheet(ws)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 42)


def top_by_reviews(rows: list[dict]) -> list[dict]:
    ranked = sorted(rows, key=lambda row: parse_int(row.get("Cantidad reviews")), reverse=True)[:20]
    return ranking_rows(ranked)


def top_by_rating(rows: list[dict]) -> list[dict]:
    ranked = sorted(rows, key=lambda row: (parse_float(row.get("Rating")), parse_int(row.get("Cantidad reviews"))), reverse=True)[:20]
    return ranking_rows(ranked)


def ranking_rows(rows: list[dict]) -> list[dict]:
    output = []
    for idx, row in enumerate(rows, 1):
        output.append({
            "Ranking": idx,
            "Nombre restaurante": row.get("Nombre restaurante", ""),
            "Comuna": row.get("Comuna", ""),
            "Región": row.get("Región", ""),
            "Rating": row.get("Rating", ""),
            "Cantidad reviews": row.get("Cantidad reviews", ""),
            "Tipo negocio": row.get("Tipo negocio", ""),
            "Nivel comercial": row.get("Nivel comercial", ""),
            "Score comercial": row.get("Score comercial", ""),
        })
    return output


def counter_rows(counter: Counter) -> list[dict]:
    return [{"Nombre": name, "Cantidad": count} for name, count in counter.most_common()]


def summary_rows(rows: list[dict]) -> list[dict]:
    return [
        {"Métrica": "Total filas", "Valor": len(rows)},
        {"Métrica": "Restaurantes válidos", "Valor": sum(1 for row in rows if row.get("Es restaurante válido") == "Sí")},
        {"Métrica": "No válidos", "Valor": sum(1 for row in rows if row.get("Es restaurante válido") == "No")},
        {"Métrica": "Alto potencial", "Valor": sum(1 for row in rows if row.get("Nivel comercial") == "Alto potencial")},
        {"Métrica": "Medio potencial", "Valor": sum(1 for row in rows if row.get("Nivel comercial") == "Medio potencial")},
        {"Métrica": "Bajo potencial", "Valor": sum(1 for row in rows if row.get("Nivel comercial") == "Bajo potencial")},
        {"Métrica": "Posibles cadenas/franquicias", "Valor": sum(1 for row in rows if row.get("Posible cadena/franquicia") == "Sí")},
        {"Métrica": "Posibles duplicados entre comunas", "Valor": sum(1 for row in rows if row.get("Posible duplicado entre comunas") == "Sí")},
    ]


def main() -> int:
    base_fields, rows = load_rows(INPUT_XLSX)
    enriched = enrich_rows(rows)
    output_fields = base_fields + [field for field in ADDED_FIELDS if field not in base_fields]

    write_csv(OUTPUT_CSV, output_fields, enriched)

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Base normalizada", output_fields, enriched)
    add_sheet(wb, "Resumen comercial", SUMMARY_FIELDS, summary_rows(enriched))
    add_sheet(wb, "Top reviews", TOP_FIELDS, top_by_reviews(enriched))
    add_sheet(wb, "Top rating", TOP_FIELDS, top_by_rating(enriched))
    add_sheet(wb, "Comunas", COUNT_FIELDS, counter_rows(Counter(row.get("Comuna") for row in enriched)))
    add_sheet(wb, "Categorías", COUNT_FIELDS, counter_rows(Counter(row.get("Categoría") for row in enriched)))
    add_sheet(wb, "Tipos negocio", COUNT_FIELDS, counter_rows(Counter(row.get("Tipo negocio") for row in enriched)))
    add_sheet(wb, "Niveles comerciales", COUNT_FIELDS, counter_rows(Counter(row.get("Nivel comercial") for row in enriched)))
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)

    print(f"Archivo Excel: {OUTPUT_XLSX}")
    print(f"Archivo CSV: {OUTPUT_CSV}")
    print(f"Filas procesadas: {len(enriched)}")
    print(f"Columnas base preservadas: {len(base_fields)}")
    print(f"Columnas finales: {len(output_fields)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

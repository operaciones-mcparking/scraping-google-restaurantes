import csv
import html
import json
import re
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "enriquecimiento_redes_sociales_piloto.json"

ADDED_FIELDS = ["Fuente Redes", "Observaciones Redes"]


def project_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def load_config() -> dict:
    config_path = Path(sys.argv[1]) if len(sys.argv) > 1 else CONFIG_PATH
    if not config_path.is_absolute():
        config_path = ROOT / config_path
    return json.loads(config_path.read_text(encoding="utf-8"))


def log_line(config: dict, message: str) -> None:
    log_path = project_path(config["log"])
    log_path.parent.mkdir(parents=True, exist_ok=True)
    line = f"{datetime.now().isoformat(timespec='seconds')} {message}"
    print(line)
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def strip_accents(text: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", text or "") if unicodedata.category(ch) != "Mn")


def norm_text(text: str) -> str:
    text = strip_accents(str(text or "")).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def choose_input(config: dict) -> Path:
    preferred = project_path(config["entradaExcelPreferida"])
    fallback = project_path(config["entradaExcelFallback"])
    if preferred.exists():
        return preferred
    if fallback.exists():
        return fallback
    raise FileNotFoundError(f"No existe entrada: {preferred} ni {fallback}")


def read_rows(config: dict) -> tuple[list[str], list[dict]]:
    path = choose_input(config)
    wb = load_workbook(path)
    ws = wb[config["hojaEntrada"]]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    max_restaurants = config.get("maxRestaurantes")
    last_row = ws.max_row if max_restaurants in (None, "", 0) else min(ws.max_row, int(max_restaurants) + 1)
    for r in range(2, last_row + 1):
        rows.append({headers[c - 1]: ws.cell(r, c).value or "" for c in range(1, ws.max_column + 1)})
    return headers, rows


def fetch(url: str, timeout: int) -> tuple[str, str]:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        content_type = response.headers.get("content-type", "")
        body = response.read(1_500_000).decode("utf-8", errors="ignore")
        return content_type, body


def absolutize(base_url: str, href: str) -> str:
    href = html.unescape(href or "").strip()
    return urllib.parse.urljoin(base_url, href)


def extract_links(base_url: str, body: str) -> list[str]:
    hrefs = re.findall(r"""href=["']([^"']+)["']""", body, flags=re.I)
    return [absolutize(base_url, href) for href in hrefs]


def clean_instagram(url: str) -> tuple[str, str]:
    parsed = urllib.parse.urlparse(url)
    host = parsed.netloc.lower().replace("www.", "")
    if host not in {"instagram.com", "m.instagram.com"}:
        return "", ""
    parts = [part for part in parsed.path.split("/") if part]
    if not parts:
        return "", ""
    username = parts[0]
    blocked = {"p", "reel", "reels", "stories", "explore", "accounts", "about", "developer"}
    if username.lower() in blocked:
        return "", ""
    return f"https://www.instagram.com/{username}/", username


def clean_facebook(url: str) -> str:
    parsed = urllib.parse.urlparse(url)
    host = parsed.netloc.lower().replace("www.", "")
    if host not in {"facebook.com", "m.facebook.com", "fb.com"}:
        return ""
    path = parsed.path.strip("/")
    blocked_prefixes = ("sharer", "plugins", "dialog", "login", "events", "groups")
    if not path or path.lower().startswith(blocked_prefixes):
        return ""
    return f"https://www.facebook.com/{path.split('?')[0].strip('/')}"


def google_search_url(query: str) -> str:
    return "https://www.google.com/search?q=" + urllib.parse.quote_plus(query)


def confidence_from_official_site(row: dict, url: str) -> str:
    if not url:
        return ""
    return "Alta"


def enrich_row(row: dict, config: dict) -> dict:
    row = dict(row)
    name = row.get("Nombre restaurante", "")
    comuna = row.get("Comuna", "")
    website = str(row.get("Sitio web") or "").strip()
    observations = []
    source = []

    row["Tiene Instagram"] = ""
    row["Calidad Instagram"] = ""
    row["Tiene Facebook"] = ""
    row["Calidad Facebook"] = ""
    row["Fuente Redes"] = ""
    row["Observaciones Redes"] = ""

    if not website:
        observations.append("Sin sitio web oficial; queda búsqueda sugerida para revisión manual.")
        observations.append(f"Instagram sugerido: {google_search_url(f'{name} {comuna} Instagram')}")
        observations.append(f"Facebook sugerido: {google_search_url(f'{name} {comuna} Facebook')}")
        row["Tiene Instagram"] = ""
        row["Tiene Facebook"] = ""
        row["_error_sitio_web"] = "No"
        row["Observaciones Redes"] = "; ".join(observations)
        return row

    direct_ig_url, direct_ig_user = clean_instagram(website)
    direct_fb_url = clean_facebook(website)
    if direct_ig_url or direct_fb_url:
        if direct_ig_url:
            row["Instagram URL"] = direct_ig_url
            row["Instagram Usuario"] = direct_ig_user
            row["Tiene Instagram"] = "Sí"
            row["Calidad Instagram"] = "Alta"
            source.append("Instagram declarado como sitio web en Google Maps")
        else:
            row["Tiene Instagram"] = ""
            observations.append(f"Instagram sugerido: {google_search_url(f'{name} {comuna} Instagram')}")

        if direct_fb_url:
            row["Facebook URL"] = direct_fb_url
            row["Tiene Facebook"] = "Sí"
            row["Calidad Facebook"] = "Alta"
            source.append("Facebook declarado como sitio web en Google Maps")
        else:
            row["Tiene Facebook"] = ""
            observations.append(f"Facebook sugerido: {google_search_url(f'{name} {comuna} Facebook')}")

        row["Fuente Redes"] = "; ".join(source)
        row["_error_sitio_web"] = "No"
        row["Observaciones Redes"] = "; ".join(observations)
        return row

    try:
        _, body = fetch(website, int(config["timeoutSegundos"]))
        links = extract_links(website, body)
    except Exception as exc:
        observations.append(f"Error accediendo sitio web oficial: {exc}")
        observations.append(f"Instagram sugerido: {google_search_url(f'{name} {comuna} Instagram')}")
        observations.append(f"Facebook sugerido: {google_search_url(f'{name} {comuna} Facebook')}")
        row["_error_sitio_web"] = "Sí"
        row["Observaciones Redes"] = "; ".join(observations)
        return row

    instagram_candidates = []
    facebook_candidates = []
    for link in links:
        ig_url, ig_user = clean_instagram(link)
        if ig_url:
            instagram_candidates.append((ig_url, ig_user))
        fb_url = clean_facebook(link)
        if fb_url:
            facebook_candidates.append(fb_url)

    if instagram_candidates:
        ig_url, ig_user = instagram_candidates[0]
        row["Instagram URL"] = ig_url
        row["Instagram Usuario"] = ig_user
        row["Tiene Instagram"] = "Sí"
        row["Calidad Instagram"] = confidence_from_official_site(row, ig_url)
        source.append("Instagram desde sitio web oficial")
    else:
        row["Tiene Instagram"] = "No"
        observations.append(f"No se encontró Instagram en sitio web. Búsqueda sugerida: {google_search_url(f'{name} {comuna} Instagram')}")

    if facebook_candidates:
        row["Facebook URL"] = facebook_candidates[0]
        row["Tiene Facebook"] = "Sí"
        row["Calidad Facebook"] = confidence_from_official_site(row, facebook_candidates[0])
        source.append("Facebook desde sitio web oficial")
    else:
        row["Tiene Facebook"] = "No"
        observations.append(f"No se encontró Facebook en sitio web. Búsqueda sugerida: {google_search_url(f'{name} {comuna} Facebook')}")

    row["Fuente Redes"] = "; ".join(source) if source else "Sitio web oficial revisado sin redes detectadas"
    row["_error_sitio_web"] = "No"
    row["Observaciones Redes"] = "; ".join(observations)
    return row


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        writer.writerows(rows)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 55)


def write_xlsx(path: Path, fields: list[str], rows: list[dict], summary: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Restaurantes"
    ws.append(fields)
    for row in rows:
        ws.append([row.get(field, "") for field in fields])
    style_sheet(ws)

    sws = wb.create_sheet("Resumen")
    sws.append(["Métrica", "Valor"])
    for item in summary:
        sws.append([item["Métrica"], item["Valor"]])
    style_sheet(sws)
    wb.save(path)


def summary_rows(rows: list[dict]) -> list[dict]:
    instagram = sum(1 for row in rows if row.get("Tiene Instagram") == "Sí")
    facebook = sum(1 for row in rows if row.get("Tiene Facebook") == "Sí")
    instagram_alta = sum(1 for row in rows if row.get("Calidad Instagram") == "Alta")
    facebook_alta = sum(1 for row in rows if row.get("Calidad Facebook") == "Alta")
    no_networks = sum(1 for row in rows if row.get("Tiene Instagram") != "Sí" and row.get("Tiene Facebook") != "Sí")
    site_errors = sum(1 for row in rows if row.get("_error_sitio_web") == "Sí")
    return [
        {"Métrica": "Filas procesadas", "Valor": len(rows)},
        {"Métrica": "Instagram encontrado", "Valor": instagram},
        {"Métrica": "Facebook encontrado", "Valor": facebook},
        {"Métrica": "Confianza Alta total", "Valor": instagram_alta + facebook_alta},
        {"Métrica": "Instagram confianza Alta", "Valor": instagram_alta},
        {"Métrica": "Facebook confianza Alta", "Valor": facebook_alta},
        {"Métrica": "Sin redes encontradas", "Valor": no_networks},
        {"Métrica": "Errores de sitio web", "Valor": site_errors},
        {"Métrica": "Búsquedas Google ejecutadas", "Valor": 0},
    ]


def main() -> int:
    config = load_config()
    base_fields, rows = read_rows(config)
    fields = base_fields + [field for field in ADDED_FIELDS if field not in base_fields]
    enriched = []

    log_line(config, f"Inicio enriquecimiento piloto. Filas: {len(rows)}")
    for index, row in enumerate(rows, 1):
        log_line(config, f"Procesando {index}/{len(rows)}: {row.get('Nombre restaurante')} | Sitio: {row.get('Sitio web')}")
        enriched.append(enrich_row(row, config))
        time.sleep(float(config["pausaSegundos"]))

    summary = summary_rows(enriched)
    public_rows = [{key: value for key, value in row.items() if not key.startswith("_")} for row in enriched]
    write_csv(project_path(config["salidaCsv"]), fields, public_rows)
    write_xlsx(project_path(config["salidaExcel"]), fields, public_rows, summary)
    project_path(config["salidaJson"]).write_text(json.dumps({"summary": summary, "rows": public_rows}, ensure_ascii=False, indent=2), encoding="utf-8")
    log_line(config, f"Excel creado: {project_path(config['salidaExcel'])}")
    log_line(config, f"CSV creado: {project_path(config['salidaCsv'])}")
    for item in summary:
        log_line(config, f"{item['Métrica']}: {item['Valor']}")
    log_line(config, "Fin enriquecimiento piloto")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

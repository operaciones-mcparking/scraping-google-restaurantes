import argparse
import csv
import json
import re
import sqlite3
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "configs" / "actualizacion_incremental.json"


def project_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def load_config(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def strip_accents(text: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", str(text or "")) if unicodedata.category(ch) != "Mn")


def norm_text(text: str) -> str:
    text = strip_accents(text).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def norm_coord(value) -> str:
    if value in (None, ""):
        return ""
    try:
        return f"{float(str(value).replace(',', '.')):.6f}"
    except ValueError:
        return norm_text(value)


def key_parts(row: dict) -> dict:
    url = str(row.get("Google Maps URL") or "").strip()
    name = norm_text(row.get("Nombre normalizado") or row.get("Nombre restaurante"))
    address = norm_text(row.get("Dirección"))
    comuna = norm_text(row.get("Comuna"))
    lat = norm_coord(row.get("Latitud"))
    lng = norm_coord(row.get("Longitud"))
    key1 = f"url::{url}" if url else ""
    key2 = f"name_address_comuna::{name}|{address}|{comuna}" if name and address and comuna else ""
    key3 = f"name_lat_lng::{name}|{lat}|{lng}" if name and lat and lng else ""
    selected = key1 or key2 or key3
    return {"key1": key1, "key2": key2, "key3": key3, "selected": selected}


def read_excel(path: Path, sheet_name: str) -> tuple[list[str], list[dict]]:
    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        rows.append({headers[c - 1]: ws.cell(r, c).value or "" for c in range(1, ws.max_column + 1)})
    return headers, rows


def read_csv(path: Path) -> tuple[list[str], list[dict]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        return headers, list(reader)


def connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def create_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS restaurantes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            unique_key TEXT NOT NULL UNIQUE,
            key_google_maps_url TEXT,
            key_nombre_direccion_comuna TEXT,
            key_nombre_lat_lng TEXT,
            nombre_restaurante TEXT,
            comuna TEXT,
            region TEXT,
            data_json TEXT NOT NULL,
            origen TEXT NOT NULL,
            fecha_insertado TEXT NOT NULL,
            fecha_actualizado TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS actualizaciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            modo TEXT NOT NULL,
            encontrados INTEGER NOT NULL,
            insertados INTEGER NOT NULL,
            duplicados INTEGER NOT NULL,
            errores INTEGER NOT NULL,
            notas TEXT
        )
        """
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_restaurantes_key1 ON restaurantes(key_google_maps_url)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_restaurantes_key2 ON restaurantes(key_nombre_direccion_comuna)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_restaurantes_key3 ON restaurantes(key_nombre_lat_lng)")
    conn.commit()


def existing_match(conn: sqlite3.Connection, keys: dict) -> bool:
    for column, value in (
        ("key_google_maps_url", keys["key1"]),
        ("key_nombre_direccion_comuna", keys["key2"]),
        ("key_nombre_lat_lng", keys["key3"]),
    ):
        if not value:
            continue
        found = conn.execute(f"SELECT 1 FROM restaurantes WHERE {column} = ? LIMIT 1", (value,)).fetchone()
        if found:
            return True
    return False


def insert_rows(conn: sqlite3.Connection, rows: list[dict], origen: str, dry_run: bool = False) -> dict:
    summary = {"encontrados": len(rows), "insertados": 0, "duplicados": 0, "errores": 0}
    now = datetime.now().isoformat(timespec="seconds")
    for row in rows:
        keys = key_parts(row)
        if not keys["selected"]:
            summary["errores"] += 1
            continue
        if existing_match(conn, keys):
            summary["duplicados"] += 1
            continue
        if dry_run:
            summary["insertados"] += 1
            continue
        conn.execute(
            """
            INSERT INTO restaurantes (
                unique_key, key_google_maps_url, key_nombre_direccion_comuna, key_nombre_lat_lng,
                nombre_restaurante, comuna, region, data_json, origen, fecha_insertado, fecha_actualizado
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                keys["selected"],
                keys["key1"],
                keys["key2"],
                keys["key3"],
                row.get("Nombre restaurante", ""),
                row.get("Comuna", ""),
                row.get("Región", ""),
                json.dumps(row, ensure_ascii=False),
                origen,
                now,
                now,
            ),
        )
        summary["insertados"] += 1
    if not dry_run:
        conn.commit()
    return summary


def all_rows(conn: sqlite3.Connection) -> list[dict]:
    rows = []
    for (data_json,) in conn.execute("SELECT data_json FROM restaurantes ORDER BY id"):
        rows.append(json.loads(data_json))
    return rows


def write_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        writer.writerows([{header: row.get(header, "") for header in headers} for row in rows])


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 45)


def write_xlsx(path: Path, headers: list[str], rows: list[dict], summary: dict | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Base restaurantes"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    style_sheet(ws)
    if summary:
        sws = wb.create_sheet("Resumen")
        sws.append(["Métrica", "Valor"])
        for key, value in summary.items():
            sws.append([key, value])
        style_sheet(sws)
    wb.save(path)


def get_headers_from_db(conn: sqlite3.Connection, fallback_headers: list[str]) -> list[str]:
    rows = all_rows(conn)
    headers = list(fallback_headers)
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    return headers


def log_update(conn: sqlite3.Connection, mode: str, summary: dict, notes: str = "") -> None:
    conn.execute(
        """
        INSERT INTO actualizaciones (fecha, modo, encontrados, insertados, duplicados, errores, notas)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            datetime.now().isoformat(timespec="seconds"),
            mode,
            summary.get("encontrados", 0),
            summary.get("insertados", 0),
            summary.get("duplicados", 0),
            summary.get("errores", 0),
            notes,
        ),
    )
    conn.commit()


def import_base(config: dict) -> dict:
    headers, rows = read_excel(project_path(config["entradaBaseExcel"]), config["hojaBaseExcel"])
    conn = connect(project_path(config["baseDatos"]))
    create_schema(conn)
    before = conn.execute("SELECT COUNT(*) FROM restaurantes").fetchone()[0]
    summary = insert_rows(conn, rows, "importacion_base")
    after = conn.execute("SELECT COUNT(*) FROM restaurantes").fetchone()[0]
    summary["total_antes"] = before
    summary["total_final"] = after
    log_update(conn, "importacion_base", summary)
    export_current(config, headers, extra_summary=summary)
    conn.close()
    return summary


def merge_file(config: dict, input_path: Path, dry_run: bool = False) -> dict:
    headers, rows = read_csv(input_path) if input_path.suffix.lower() == ".csv" else read_excel(input_path, "Restaurantes")
    conn = connect(project_path(config["baseDatos"]))
    create_schema(conn)
    before = conn.execute("SELECT COUNT(*) FROM restaurantes").fetchone()[0]
    summary = insert_rows(conn, rows, f"incremental:{input_path.name}", dry_run=dry_run)
    after = conn.execute("SELECT COUNT(*) FROM restaurantes").fetchone()[0]
    summary["total_antes"] = before
    summary["total_final"] = after if not dry_run else before
    if not dry_run:
        log_update(conn, "incremental", summary, input_path.name)
        export_current(config, get_headers_from_db(conn, headers), extra_summary=summary)
    conn.close()
    return summary


def export_current(config: dict, headers: list[str] | None = None, extra_summary: dict | None = None) -> dict:
    conn = connect(project_path(config["baseDatos"]))
    create_schema(conn)
    rows = all_rows(conn)
    headers = get_headers_from_db(conn, headers or [])
    summary = {
        "Total base final": len(rows),
        "Archivo Excel generado": str(project_path(config["excelSalida"])),
    }
    if extra_summary:
        summary = {**extra_summary, **summary}
    write_csv(project_path(config["csvSalida"]), headers, rows)
    write_xlsx(project_path(config["excelSalida"]), headers, rows, summary)
    conn.close()
    return summary


def dry_run_existing(config: dict) -> dict:
    headers, rows = read_excel(project_path(config["entradaBaseExcel"]), config["hojaBaseExcel"])
    sample = rows[:10]
    conn = connect(project_path(config["baseDatos"]))
    create_schema(conn)
    before = conn.execute("SELECT COUNT(*) FROM restaurantes").fetchone()[0]
    summary = insert_rows(conn, sample, "prueba_sin_scraping", dry_run=True)
    summary["total_antes"] = before
    summary["total_final"] = before
    conn.close()
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=str(DEFAULT_CONFIG))
    parser.add_argument("--import-base", action="store_true")
    parser.add_argument("--merge-file")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--export", action="store_true")
    parser.add_argument("--test-no-scraping", action="store_true")
    args = parser.parse_args()

    config = load_config(Path(args.config))
    if args.import_base:
        summary = import_base(config)
    elif args.merge_file:
        summary = merge_file(config, project_path(args.merge_file), dry_run=args.dry_run)
    elif args.export:
        summary = export_current(config)
    elif args.test_no_scraping:
        summary = dry_run_existing(config)
    else:
        parser.error("Indica --import-base, --merge-file, --export o --test-no-scraping")

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
